"""母版处理（完整版手册 §7 / §14）。

母版（Word 模板）只提供**结构与槽位**：样式 id、``outlineLvl``、多级编号骨架、
页眉页脚占位。**一切格式值由引擎从 config 无条件写入**（手册 §5 铁律 2），
所以母版里的值只是「默认值」，会被 config 覆盖。

配置口径（手册 §6.1）：点号键，如 ``font.name``、``font.size``、``font.bold``、
``line_spacing``；引擎按此读取，不再兼容下划线扁平键。

W0 阶段用 python-docx 自带的默认模板作为最小母版，先跑通链路；
真实 ``.dotx`` 解包入库（含 ``numbering.xml`` 多级编号、``TOC 1..9``）在后续阶段替换。
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document as _Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..config import Config, build_style_table
from ..units import to_points
from .ooxml import set_style_outline_level

__all__ = [
    "MONO_EAST_ASIA",
    "MONO_FONT",
    "TEMPLATE_VERSION",
    "create_document",
    "ensure_font_table",
    "heading_indent",
    "heading_numbering_enabled",
    "iter_styles",
]

#: 母版版本：随解包入库的模板一起提交，CI 可校验
TEMPLATE_VERSION = "0.1"

#: 必须存在的样式槽位（手册 §14.1）：母版缺任何一个都算不完整
REQUIRED_STYLES = (
    "Normal",
    "Heading1",
    "Heading2",
    "Heading3",
    "BodyText",
    "Caption",
    "TOC1",
    "TOC9",
)

#: 标题样式名 -> python-docx 样式键
HEADING_STYLE_NAMES = tuple(f"Heading{level}" for level in range(1, 10))


def create_document(config: Config, template: Path | None = None) -> Any:
    """创建 docx 文档：加载母版（或引擎默认母版）+ 补齐槽位 + 写入 ``[document]``。

    **构建路径与母版构建路径共用同一套槽位补齐**：否则用引擎默认母版构建出来的
    文档会缺少 ``TOC 1..9`` 样式与多级标题编号（题注样式同理），
    而这两样正是 W3/W5 依赖的。
    """
    if template is not None:
        document = _Document(str(template))
    else:
        document = _Document()
        _extend_default_template(document, config)
    apply_document_defaults(document, config)
    return document


def heading_numbering_enabled(config: Config) -> bool:
    """``[document] heading_numbering``：标题多级编号开关，缺省开启（手册 §14 要求）。

    支持 ``heading_numbering = false`` 与 ``heading_numbering.enabled = false`` 两种写法。
    """
    value = config.setting("heading_numbering", True)
    if isinstance(value, dict):
        value = value.get("enabled", True)
    return bool(value)


def heading_indent(config: Config) -> dict[str, Any] | None:
    """``[document] heading_indent``：标题缩进（传给编号定义）。

    三态语义：

    - **不写** → ``None``（默认 ``left=720/hanging=360``，Word 常见观感）
    - ``heading_indent = false`` 或 ``left = 0`` → ``{}``（**取消缩进**，标题齐头）
    - ``[document.heading_indent] left = "1cm"`` 或点号键
      ``heading_indent.left = "1cm"`` → 自定义距离

    支持子表、点号键、标量三种写法（``Config.setting`` 只管单点路径，
    所以这里自己收拢 ``heading_indent.*`` 前缀的键）。
    """
    direct = config.document.get("heading_indent")
    nested = config.setting("heading_indent")
    spec: dict[str, Any] = {}
    if isinstance(direct, dict):
        spec.update(direct)
    if isinstance(nested, dict):
        spec.update(nested)
    for key, value in config.document.items():
        prefix = "heading_indent."
        if key.startswith(prefix):
            spec[key[len(prefix) :]] = value

    if not spec:
        if direct is None:
            return None
        if direct is False:
            return {}
        if isinstance(direct, (str, int, float)):
            return {"left": direct}
        return None
    return spec


def _extend_default_template(document: Any, config: Config | None = None) -> None:
    """补齐 §14.1 槽位（幂等）：目录样式、题注样式、字体表、多级编号、updateFields。"""
    from .template_build import add_heading_numbering, add_toc_styles, ensure_update_fields

    _drop_macro_styles(document)
    drop_styles_with_effects(document)
    add_toc_styles(document)
    _ensure_caption_styles(document)
    ensure_font_table(document)
    enabled = heading_numbering_enabled(config) if config is not None else True
    indent = heading_indent(config) if config is not None else None
    add_heading_numbering(document, enabled=enabled, indent=indent)
    ensure_update_fields(document)


#: 默认模板里带点阵 ``Courier`` 的宏样式：内容文档用不到，且该字体在 Linux 上必然缺失，
#: 留着会让 Word/LibreOffice 报「缺少字体」
UNUSED_STYLE_NAMES = frozenset({"MacroText", "MacroTextChar", "macro", "Macro Text Char"})


def _drop_macro_styles(document: Any) -> list[str]:
    """删掉用不到的宏样式，避免凭空引入 ``Courier`` 字体依赖。"""
    removed: list[str] = []
    for style in list(document.styles):
        if style.name in UNUSED_STYLE_NAMES or style.style_id in UNUSED_STYLE_NAMES:
            style.element.getparent().remove(style.element)
            removed.append(style.name)
    return removed


def _ensure_caption_styles(document: Any) -> None:
    """题注/脚注样式必须存在（``SEQ`` 域与脚注引用都依赖它们）。

    同时把母版自带的 ``Caption`` **中性化**：python-docx 默认母版的题注是
    「加粗 + 主题蓝 ``4F81BD`` + 9pt」，与「config 是唯一真源」冲突——
    没声明 ``[styles.caption]`` 时，题注不该自带一套没人要的强调色。
    清掉颜色/加粗/字号后它继承 ``Normal``（= 正文外观），config 再覆盖。
    """
    for name in ("Caption", "FootnoteText", "FootnoteReference"):
        if any(style.name == name for style in document.styles):
            continue
        style = document.styles.add_style(name, 1)  # 1 = paragraph style
        style.base_style = document.styles["Normal"]
        style.quick_style = True
        style.hidden = False
    _neutralise_caption_style(document)


def _neutralise_caption_style(document: Any) -> None:
    """去掉母版 ``Caption`` 的强调外观（颜色/加粗/字号），让它跟正文一致。"""
    from docx.oxml.ns import qn

    for style in document.styles:
        if style.name != "Caption":
            continue
        rpr = style.element.find(qn("w:rPr"))
        if rpr is None:
            return
        for name in ("w:b", "w:bCs", "w:color", "w:sz", "w:szCs"):
            for element in rpr.findall(qn(name)):
                rpr.remove(element)
        return


#: 代码块等场景的等宽字体（与 docx.py 的 DEFAULT_MONO_FONT 一致）
MONO_FONT = "Consolas"
#: 代码块里中文的回退字体
MONO_EAST_ASIA = "等线"


def drop_styles_with_effects(document: Any) -> bool:
    """断开并删除 ``word/stylesWithEffects.xml``（Word 2007 兼容部件）。

    内容文档用不到它，但它完整复制了一份样式表——里面就有 ``Courier``，
    会让我们在删掉 ``MacroText`` 之后**依然**报「缺少字体」。
    删完必须同时清掉关系与 ``[Content_Types].xml`` 里对应的 Override。
    """
    target_name = ""
    for part in list(document.part.package.parts):
        if "stylesWithEffects" in str(part.partname):
            target_name = str(part.partname)
            break
    if not target_name:
        return False
    for rel in list(document.part.rels.values()):
        if "stylesWithEffects" in getattr(rel, "reltype", ""):
            del document.part.rels[rel.rId]
    document._docgen_drop_part = target_name
    return True


def _strip_effects_part(data: bytes, part_name: str) -> bytes:
    """把已断开关系的部件从 zip 与 ``[Content_Types].xml`` 里真正删掉。"""
    import re
    import zipfile

    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename == part_name.lstrip("/"):
                continue
            payload = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = payload.decode("utf-8")
                text = re.sub(rf'<Override PartName="{re.escape(part_name)}"[^/]*/>', "", text)
                payload = text.encode("utf-8")
            target.writestr(item, payload)
    source.close()
    return buffer.getvalue()


#: 名称需要替代字体时给出的别名（``Courier`` 已随宏样式一起删除，不需要别名）
FONT_ALIASES: dict[str, str] = {}


def ensure_font_table(document: Any, extra_fonts: list[str] | None = None) -> list[str]:
    """把用到的字体登记进 ``word/fontTable.xml``，返回本次新增的字体名。

    Word 只认 fontTable / theme 里声明过的字体：缺登记会提示「缺少字体」并替换字形
    （DSH 预览侧的 "Missing fonts" 就是这里来的）。python-docx 没有 fontTable 的
    专用 API，所以直接改这个部件的 XML。
    """
    part = _font_table_part(document)
    if part is None:
        return []

    from lxml import etree

    root = etree.fromstring(part.blob)
    existing = {
        font.get(qn("w:name")) for font in root.findall(qn("w:font")) if font.get(qn("w:name"))
    }

    wanted = [MONO_FONT, MONO_EAST_ASIA, *FONT_ALIASES, *(extra_fonts or [])]
    added: list[str] = []
    for name in wanted:
        if not name or name in existing:
            continue
        root.append(_font_element(name, FONT_ALIASES.get(name)))
        existing.add(name)
        added.append(name)

    if added:
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return added


def _font_table_part(document: Any) -> Any:
    """在包部件里找 ``/word/fontTable.xml``。"""
    for part in document.part.package.parts:
        if str(part.partname).endswith("fontTable.xml"):
            return part
    return None


def _font_element(name: str, substitute: str | None = None) -> Any:
    """构造 ``<w:font w:name="X">``；中文额外写 ``w:charset="86"``（GB2312）。"""
    font = OxmlElement("w:font")
    font.set(qn("w:name"), name)
    if substitute:
        alternative = OxmlElement("w:altName")
        alternative.set(qn("w:val"), substitute)
        font.append(alternative)
    if _is_cjk(name):
        charset = OxmlElement("w:charset")
        charset.set(qn("w:val"), "86")
        font.append(charset)
    family = OxmlElement("w:family")
    family.set(qn("w:val"), "auto")
    font.append(family)
    pitch = OxmlElement("w:pitch")
    pitch.set(qn("w:val"), "variable")
    font.append(pitch)
    return font


def _is_cjk(name: str) -> bool:
    return any("\u4e00" <= char <= "\u9fff" for char in name)


def apply_document_defaults(document: Any, config: Config) -> None:
    """把 config 的字体/字号/行距写进**载体样式**（正文继承它们）。

    两条映射（语义样式 → Word 内置样式）：

    | config | Word 样式 | 因为 |
    |---|---|---|
    | ``[document]`` + ``[styles.body]`` | ``Normal`` | 正文段落默认继承它 |
    | ``[styles.caption]`` | ``Caption`` | 题注段落用 ``style="Caption"`` |

    为什么必须这么写：正文/题注段落本身只带段落级属性（``line_spacing`` /
    ``first_line_chars``），``font.color`` 这类字符级属性没有 run 去承载。
    只把值读进 config 而没人写进样式，就是「config 写了但不生效」——
    本项目最隐蔽的失效方式（实测踩到：``[styles.caption]`` 完全没作用）。

    取值**必须读 ``build_style_table`` 的展开结果**：``config.styles`` 里还留着
    未解析的 ``@ref``（如 ``@document.font.size``），直接用会拿到字符串，
    然后在单位换算里炸掉。展开表也正是渲染层 ``state.styles`` 用的那一份。
    """
    table = build_style_table(config)
    resolved_doc = table.get("document", {})
    doc_settings: dict[str, Any] = resolved_doc if isinstance(resolved_doc, dict) else {}
    raw_body = table.get("body", {})
    body: dict[str, Any] = raw_body if isinstance(raw_body, dict) else {}

    normal = _style_by_id(document, "Normal")
    # body 覆盖 document（样式比文档默认更具体）
    _apply_font_properties(normal, doc_settings)
    _apply_font_properties(normal, body)
    line_spacing = body.get("line_spacing", doc_settings.get("line_spacing"))
    if line_spacing is not None:
        normal.paragraph_format.line_spacing = float(line_spacing)

    # 题注：段落用的是内置 Caption 样式，所以把 config 写在样式上
    raw_caption = table.get("caption", {})
    caption = raw_caption if isinstance(raw_caption, dict) else {}
    if caption:
        caption_style = _style_or_none(document, "Caption")
        if caption_style is not None:
            _apply_font_properties(caption_style, caption)

    for level in range(1, 10):
        style = _style_by_id(document, f"Heading{level}")
        set_style_outline_level(style, level - 1)
        _apply_style_indent(style, config, level)
        # 标题的**样式定义**也要吃 config：只写 run 级的话，Word 样式面板与
        # 「新敲一个标题」仍是母版的主题蓝（实测反馈），与「config 唯一真源」矛盾。
        heading_props = table.get(f"heading{level}")
        if isinstance(heading_props, dict) and heading_props:
            _apply_font_properties(style, heading_props)


def _apply_font_properties(style: Any, properties: dict[str, Any]) -> None:
    """把一组（已展开的）样式属性里的**字体类**值写进 Word 样式。

    同时接受点号键展开出的嵌套表（``{"font": {...}}``）与扁平键（``font.name``），
    与渲染层 :func:`docx._apply_run_style` 的取值口径一致。
    """
    nested = properties.get("font")
    nested = nested if isinstance(nested, dict) else {}

    def pick(key: str) -> Any:
        if key in nested:
            return nested[key]
        return properties.get(f"font.{key}")

    name = pick("name")
    size = pick("size")
    color = pick("color")
    bold = pick("bold")
    italic = pick("italic")
    underline = pick("underline")
    strike = pick("strike")

    if name:
        _set_style_fonts(style, str(name))
    if size is not None:
        style.font.size = to_points(size)
        # szCs：复杂文种字号。只写 w:sz 时中文在部分场景不跟随
        from .ooxml import insert_child

        points = to_points(size)
        insert_child(style.element.get_or_add_rPr(), "w:szCs", **{"w:val": str(int(points.pt * 2))})
    if color:
        style.font.color.rgb = _rgb(str(color))
        # 主题色优先级高于显式值：不清掉 themeColor，颜色等于没设
        color_element = style.element.find(f"{qn('w:rPr')}/{qn('w:color')}")
        if color_element is not None:
            for theme in ("w:themeColor", "w:themeShade", "w:themeTint"):
                color_element.attrib.pop(qn(theme), None)
    if bold is not None:
        style.font.bold = bool(bold)
    if italic is not None:
        style.font.italic = bool(italic)
    if underline is not None:
        style.font.underline = bool(underline)
    if strike is not None:
        style.font.strike = bool(strike)


def _style_or_none(document: Any, style_id: str) -> Any:
    """按 id 取样式；不存在返回 ``None``（自定义母版可能没有该槽位）。"""
    try:
        return _style_by_id(document, style_id)
    except KeyError:  # pragma: no cover - 依赖 python-docx 的异常类型
        return None


def _rgb(value: str) -> Any:
    """``color`` 写 ``111111`` 或 ``#111111`` 都接受（OOXML 里不带 ``#``）。"""
    from docx.shared import RGBColor

    return RGBColor.from_string(value.lstrip("#").upper())


def _apply_style_indent(style: Any, config: Config, level: int) -> None:
    """标题样式的缩进：``[document] heading_indent`` 打底，``[styles.headingN.indent]` 覆盖。

    写在**样式**上（而非只写编号定义）的原因：样式级 ``w:ind`` 优先于编号的 ``w:pPr``，
    这样「取消缩进」不依赖编号定义是否被 Word 采纳。
    """
    from docx.oxml.ns import qn

    from .ooxml import insert_child
    from .template_build import _indent_element

    indent = heading_indent(config)
    style_spec = config.styles.get(f"heading{level}")
    if style_spec is not None:
        own = style_spec.properties.get("indent")
        if isinstance(own, dict):
            indent = dict(own)
        elif own is False:
            indent = {}
    element = _indent_element(indent)
    ppr = style.element.get_or_add_pPr()
    for existing in ppr.findall(qn("w:ind")):
        ppr.remove(existing)
    if element is None:
        return
    # lxml 的属性键是完整 URI（``{uri}left``），insert_child 需要 ``w:left`` 形式
    attributes = {f"w:{name.split('}')[-1]}": value for name, value in element.attrib.items()}
    insert_child(ppr, "w:ind", **attributes)


def _style_by_id(document: Any, style_id: str) -> Any:
    """按 ``style_id`` 取样式。

    python-docx 1.2 起「按 style_id 查找」已弃用，要按**样式名**查找；
    这里通过 ``style_id -> style.name`` 反查一次，既明确又避免弃用警告。
    """
    styles = document.styles
    for style in styles:
        if style.style_id == style_id:
            return styles[style.name]
    raise KeyError(f"母版缺少样式: {style_id}")


def _set_style_fonts(style: Any, name: str) -> None:
    """样式级字体：四个 ``rFonts`` 属性一起写（缺 eastAsia 中文会回退）。

    还要**清掉主题字体属性**（``w:asciiTheme`` 等）：OOXML 里主题属性优先于
    显式值，留着它们会让 ``font.name = "黑体"`` 只对外文生效、
    样式面板也仍显示主题字体（实测反馈）。
    """
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(theme), None)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), name)


def iter_styles(document: Any) -> list[tuple[str, str]]:
    """列出文档里的样式 ``(style_id, name)``，供母版校准与探针使用。"""
    return [(style.style_id, style.name) for style in document.styles]
