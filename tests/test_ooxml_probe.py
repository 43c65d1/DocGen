"""OOXML 语义探针（完整版手册 §10.1）。

只断言 Word 专有结构：样式引用、``outlineLvl``、``rFonts`` 四属性、
``updateFields``、分页符、表格边框。**不做全文 document.xml 快照**（脆弱）。
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document

from docgen.config import Config, load_config
from docgen.ir import BlockNode
from docgen.parser.dfm import parse_dfm
from docgen.render.docx import render_docx_bytes

FIXTURES = Path(__file__).parent / "fixtures"

CONFIG = """\
config_version = "0.1"

[document]
font.name = "宋体"
font.size = "12pt"
line_spacing = 1.5

[styles.heading1]
font.name = "黑体"
font.size = "18pt"
font.bold = true
font.color = "1F3864"

[styles.strong]
font.bold = true

[styles.keyword]
font.color = "0000FF"
font.bold = true

[styles.body]
line_spacing = 1.5
first_line_chars = 200

[code_block]
font.name = "Consolas"
font.size = "10pt"
background = "F5F5F5"
border = "CCCCCC"
"""


def _config(tmp_path: Path) -> Config:
    path = tmp_path / "config.toml"
    path.write_text(CONFIG, encoding="utf-8")
    return load_config(path)


def _build(
    markdown: str, tmp_path: Path, base_dir: Path | None = None
) -> tuple[bytes, str, str, str]:
    """返回 (docx 字节, document.xml, styles.xml, settings.xml)。"""
    blocks: list[BlockNode] = parse_dfm(markdown)
    data = render_docx_bytes(blocks, _config(tmp_path), base_dir=base_dir)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
    return data, document_xml, styles_xml, settings_xml


# --------------------------------------------------------------------------
# 母版 / 样式槽位
# --------------------------------------------------------------------------
def _styles_xml(markdown: str, tmp_path: Path) -> str:
    """只要 styles.xml（母版槽位相关断言用）。"""
    return _build(markdown, tmp_path)[2]


def _document_xml(markdown: str, tmp_path: Path) -> str:
    """只要 document.xml（正文结构相关断言用）。"""
    return _build(markdown, tmp_path)[1]


def test_styles_have_required_ids(tmp_path: Path) -> None:
    blocks = parse_dfm("# 标题\n")
    data = render_docx_bytes(blocks, _config(tmp_path))
    document = Document(io.BytesIO(data))
    style_ids = {style.style_id for style in document.styles}
    for required in ("Normal", "Heading1", "Heading2", "Heading3"):
        assert required in style_ids, required


def test_headings_carry_outline_level(tmp_path: Path) -> None:
    """标题必须带 w:outlineLvl，否则导航窗格与 TOC 不认。"""
    styles = _styles_xml("# 一级\n\n## 二级\n", tmp_path)
    assert "outlineLvl" in styles
    assert 'w:val="0"' in styles


def test_normal_style_has_four_font_attributes(tmp_path: Path) -> None:
    """中英混排必须四个 rFonts 属性齐全。"""
    styles = _styles_xml("段落\n", tmp_path)
    assert 'w:ascii="宋体"' in styles
    assert 'w:hAnsi="宋体"' in styles
    assert 'w:eastAsia="宋体"' in styles
    assert 'w:cs="宋体"' in styles


def test_settings_enable_update_fields(tmp_path: Path) -> None:
    """updateFields 让 Word 打开时自动更新目录/编号域。"""
    settings = _build("# 标题\n", tmp_path)[3]
    assert "updateFields" in settings


# --------------------------------------------------------------------------
# 块级结构
# --------------------------------------------------------------------------
def test_heading_uses_heading_style(tmp_path: Path) -> None:
    document = _document_xml("# 一级\n\n## 二级\n", tmp_path)
    assert 'w:val="Heading1"' in document
    assert 'w:val="Heading2"' in document


def test_paragraph_text_is_present(tmp_path: Path) -> None:
    document = _document_xml("正文段落\n", tmp_path)
    assert "正文段落" in document


def test_first_line_indent_uses_chars(tmp_path: Path) -> None:
    """首行缩进必须用 ``w:ind@w:firstLineChars``（字符单位），不用 ``w:firstLine``。

    注意：``firstLineChars`` 是 ``w:ind`` 的**属性**，在 schema 里没有同名子元素；
    写成 ``<w:firstLineChars/>`` 会被 Word 忽略。
    """
    document = _document_xml("正文段落\n", tmp_path)
    assert '<w:ind w:firstLineChars="200"/>' in document
    assert "firstLine=" not in document


def test_page_break_marker(tmp_path: Path) -> None:
    document = _document_xml("[[pagebreak]]\n", tmp_path)
    assert 'w:type="page"' in document


def test_list_styles(tmp_path: Path) -> None:
    document = _document_xml("- 无序\n\n1. 有序\n", tmp_path)
    assert "ListBullet" in document
    assert "ListNumber" in document


def test_table_has_borders_and_header(tmp_path: Path) -> None:
    document = _document_xml("| 列 A | 列 B |\n| --- | --- |\n| a1 | b1 |\n", tmp_path)
    assert "tblBorders" in document
    assert "列 A" in document and "a1" in document
    assert "<w:b/>" in document or 'w:b w:val="true"' in document or "<w:b " in document


def test_code_block_uses_mono_font(tmp_path: Path) -> None:
    document = _document_xml("```python\nprint(1)\n```\n", tmp_path)
    assert "Consolas" in document
    assert "print(1)" in document


def test_code_block_background_and_border_from_config(tmp_path: Path) -> None:
    """代码块背景/边框来自 config，不是引擎写死的。"""
    document = _document_xml("```python\nprint(1)\n```\n", tmp_path)
    assert '<w:shd w:val="clear" w:color="auto" w:fill="F5F5F5"/>' in document
    assert "<w:pBdr>" in document
    assert 'w:color="CCCCCC"' in document


# --------------------------------------------------------------------------
# 行内样式：config 驱动
# --------------------------------------------------------------------------
def test_styled_run_writes_config_values(tmp_path: Path) -> None:
    """不变量：输出里的颜色/加粗就是 config 的声明值。"""
    document = _document_xml("[关键字]{.keyword}\n", tmp_path)
    assert 'w:val="0000FF"' in document
    assert "关键字" in document
    assert "<w:b/>" in document or 'w:b w:val="true"' in document or "<w:b " in document


def test_heading_run_uses_heading1_values(tmp_path: Path) -> None:
    document = _document_xml("# 标题\n", tmp_path)
    assert "黑体" in document
    assert 'w:val="1F3864"' in document
    assert "36" in document  # 18pt => half-points 36


def test_strong_marker_renders_bold(tmp_path: Path) -> None:
    """**粗体** 走 strong 样式：加粗来自 config，不是引擎写死的。"""
    data, _, _, _ = _build("**粗体**\n", tmp_path)
    document = Document(io.BytesIO(data))
    paragraph = document.paragraphs[0]
    assert paragraph.runs
    assert any(run.bold for run in paragraph.runs)


def test_strong_without_config_still_keeps_text(tmp_path: Path) -> None:
    """config 没声明 strong 时不写死加粗，但内容不能丢（config 是唯一真源）。"""
    blocks = parse_dfm("**粗体**\n")
    data = render_docx_bytes(blocks, load_config(FIXTURES / "minimal.toml"))
    document = Document(io.BytesIO(data))
    assert "粗体" in document.paragraphs[0].text


def test_superscript_and_subscript_rendering(tmp_path: Path) -> None:
    """手册 §6.2.3：H~2~O 与 E=mc^2^ 要真的变成 Word 上下标。"""
    document = _document_xml("H~2~O 与 E=mc^2^\n", tmp_path)
    assert 'w:val="superscript"' in document
    assert 'w:val="subscript"' in document
    assert "H" in document and "O" in document


def test_attribute_sup_marker_rendering(tmp_path: Path) -> None:
    document = _document_xml("[上标]{.sup}\n", tmp_path)
    assert 'w:val="superscript"' in document


def test_begin_code_block_renders_like_fence(tmp_path: Path) -> None:
    """\\begin{code:lang} 与 ```lang 产生同样的 Word 内容与等宽字体。"""
    fenced = _document_xml("```python\nprint(1)\n```\n", tmp_path)
    begun = _document_xml("\\begin{code:python}\nprint(1)\n\\end{code:python}\n", tmp_path)
    assert "print(1)" in begun
    assert "Consolas" in begun
    assert fenced.count("print(1)") == begun.count("print(1)")


def test_caption_without_label_keeps_content(tmp_path: Path) -> None:
    """没有 ``#id`` 的 caption 是普通命名块：内容照常渲染，内容不能丢。"""
    document = _document_xml("\\begin{caption}\n图 1 系统架构\n\\end{caption}\n", tmp_path)
    assert "图 1 系统架构" in document


def test_formula_renders_as_omml(tmp_path: Path) -> None:
    """块级公式应写成可编辑的 OMML（m:oMath），而不是纯文本。"""
    from docgen.render.math import find_xsl

    document = _document_xml("$$\nE = mc^2\n$$\n", tmp_path)
    if find_xsl() is not None:
        # 注意别拿 "mc" 当 OMML 的判据——公式内容里本来就有 ``E = mc^2``（原来是假阳性）
        assert "oMath" in document
        assert "<m:oMath" in document
        assert 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"' in document
    else:
        assert "E = mc^2" in document  # 无 XSL 时降级成 LaTeX 文本


def test_inline_math_renders_in_paragraph(tmp_path: Path) -> None:
    from docgen.render.math import find_xsl

    document = _document_xml("行内 $x^2$ 公式\n", tmp_path)
    assert "行内" in document
    if find_xsl() is not None:
        assert "oMath" in document


def test_plain_engine_keeps_latex_text(tmp_path: Path) -> None:
    """``[math] engine = "plain"`` 时不写 OMML，但内容不能丢。"""
    config_path = tmp_path / "config.toml"
    config_path.write_text(CONFIG + '\n[math]\nengine = "plain"\n', encoding="utf-8")
    blocks = parse_dfm("$$\n\\frac{1}{2}\n$$\n")
    data = render_docx_bytes(blocks, load_config(config_path))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        document = archive.read("word/document.xml").decode("utf-8")
    assert "oMath" not in document
    assert "frac" in document


def test_math_error_mode_raises(tmp_path: Path) -> None:
    """``on_fail = "error"`` 且公式非法时报 E5002。"""
    from docgen import errors
    from docgen.render.math import find_xsl

    if find_xsl() is None:
        pytest.skip("本机没有 MML2OMML.XSL")
    config_path = tmp_path / "config.toml"
    config_path.write_text(CONFIG + '\n[math]\non_fail = "error"\n', encoding="utf-8")
    blocks = parse_dfm("$$\n\\frac{1}{\n$$\n")
    with pytest.raises(errors.DocGenError) as caught:
        render_docx_bytes(blocks, load_config(config_path))
    assert caught.value.code == errors.E_MATH_CONVERSION


def test_caption_renders_with_seq_field_and_bookmark(tmp_path: Path) -> None:
    """题注要带书签 + SEQ 编号域 + Caption 样式（Word 打开自动编号）。"""
    document = _document_xml("\\begin{caption #fig:1}\n图 1 系统架构\n\\end{caption}\n", tmp_path)
    assert "Caption" in document
    assert "bookmarkStart" in document
    assert 'w:name="fig:1"' in document
    assert "SEQ" in document
    assert "图 " in document


def test_table_caption_uses_table_prefix(tmp_path: Path) -> None:
    document = _document_xml(
        "\\begin{caption kind=table #tab:1}\n表 1 参数\n\\end{caption}\n", tmp_path
    )
    assert "SEQ 表" in document


def test_reference_renders_as_ref_field(tmp_path: Path) -> None:
    from docgen.render.math import find_xsl  # noqa: F401  (保持导入顺序稳定)

    document = _document_xml(
        "\\begin{caption #fig:1}\n图 1\n\\end{caption}\n\n见图 [[ref:fig:1]]。\n", tmp_path
    )
    assert "REF fig:1" in document
    assert "见图" in document


def test_toc_marker_renders_as_toc_field(tmp_path: Path) -> None:
    document = _document_xml("[[toc]]\n\n# 标题\n", tmp_path)
    assert "TOC" in document
    settings = _build("[[toc]]\n", tmp_path)[3]
    assert "updateFields" in settings  # 打开时自动更新目录


def test_output_is_deterministic(tmp_path: Path) -> None:
    """确定性：同一份 IR + config 渲染两次，document.xml 必须逐字节一致。

    比 document.xml 而不是整个 docx 字节——zip 会写入时间戳，字节本身不可比。
    """
    blocks = parse_dfm("# 标题\n\n段落 **粗**\n")
    config = _config(tmp_path)
    first = _document_xml_from(render_docx_bytes(blocks, config))
    second = _document_xml_from(render_docx_bytes(blocks, config))
    assert first == second


def _document_xml_from(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


# --------------------------------------------------------------------------
# 图片
# --------------------------------------------------------------------------
def test_missing_image_is_reported_in_text(tmp_path: Path) -> None:
    document = _document_xml("[[img:pics/not-there.png]]\n", tmp_path)
    assert "图片缺失" in document


def test_existing_image_is_embedded(tmp_path: Path) -> None:
    blocks = parse_dfm("[[img:pics/photo.png]]\n")
    data = render_docx_bytes(blocks, _config(tmp_path), base_dir=FIXTURES)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert media, "图片应写入 word/media/"
